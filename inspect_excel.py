import openpyxl

wb = openpyxl.load_workbook("carga psico.xlsx")
sheet = wb.active
print("Sheet Title:", sheet.title)
print("Max Row:", sheet.max_row)
print("Max Column:", sheet.max_column)

print("\nHeaders:")
headers = [sheet.cell(1, col).value for col in range(1, sheet.max_column + 1)]
for i, h in enumerate(headers, start=1):
    print(f"Col {i}: {h}")

print("\nFirst 3 rows of data:")
for r in range(2, min(5, sheet.max_row + 1)):
    row_vals = [sheet.cell(r, col).value for col in range(1, sheet.max_column + 1)]
    print(f"Row {r}: {row_vals}")
