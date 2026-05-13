#!/usr/bin/env python3
"""
Script para generar Excel de prueba de docentes
Formato: Nombre | Apellido | Email
"""

try:
    from openpyxl import Workbook
except ImportError:
    print("Instalando openpyxl...")
    import subprocess
    subprocess.check_call(["pip", "install", "openpyxl"])
    from openpyxl import Workbook

# Crear libro
wb = Workbook()
ws = wb.active
ws.title = "Docentes"

# Agregar header
ws['A1'] = 'Nombre'
ws['B1'] = 'Apellido'
ws['C1'] = 'Email'

# Datos de prueba - docentes variados
docentes = [
    ("Juan", "Pérez", "juan.perez@upeu.edu.pe"),
    ("María", "García López", "maria.garcia@upeu.edu.pe"),
    ("Juan Tomás", "Pérez Del Aguila", "juan.tomás@upeu.edu.pe"),
    ("Carlos", "Rodríguez", "carlos.rodriguez@upeu.edu.pe"),
    ("Ana", "Martínez Gómez", "ana.martinez@upeu.edu.pe"),
    ("Roberto", "López Flores", "roberto.lopez@upeu.edu.pe"),
    ("Sofia", "Díaz", "sofia.diaz@upeu.edu.pe"),
    ("Luis", "Fernández", "luis.fernandez@upeu.edu.pe"),
    ("Patricia", "Sánchez", "patricia.sanchez@upeu.edu.pe"),
    ("Miguel", "Quispe Condori", "miguel.quispe@upeu.edu.pe"),
]

# Insertar datos
for idx, (nombre, apellido, email) in enumerate(docentes, start=2):
    ws[f'A{idx}'] = nombre
    ws[f'B{idx}'] = apellido
    ws[f'C{idx}'] = email

# Ajustar ancho de columnas
ws.column_dimensions['A'].width = 15
ws.column_dimensions['B'].width = 20
ws.column_dimensions['C'].width = 30

# Guardar archivo
output_path = "docentes_prueba.xlsx"
wb.save(output_path)
print(f"✓ Archivo creado: {output_path}")
print(f"✓ Total de docentes: {len(docentes)}")
print("\nDocentes incluidos:")
for nombre, apellido, email in docentes:
    print(f"  - {nombre} {apellido} ({email})")
